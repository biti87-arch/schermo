"""Controlla sorgenti/meteo.json sulle «Tabelle del Tempo Atmosferico v1.0».
Uso: python3 sorgenti/controlla_meteo.py [percorso del .docx]   (stampa solo errori e conteggi)
Legge il Word con python-docx (non l'estratto usato da crea_meteo.py), quindi è un controllo indipendente:
1. ogni tabella del tempo copre 01–00 senza buchi né sovrapposizioni ed è uguale a quella del Word;
2. ogni voce di ogni riga (e ogni condizione del Sotterraneo) ha il suo effetto;
3. i testi (effetti, zone, particolarità, regole, Foraggiare, acqua, Fame e sete) sono uguali a paragrafi o celle del Word;
4. le regole che il motore applica (durate, varianti, modificatori di Foraggiare, moltiplicatori dell'acqua) hanno la loro frase nel Word."""
import json, os, re, sys
import docx
R = os.path.dirname(os.path.abspath(__file__))
M = json.load(open(os.path.join(R, 'meteo.json'), encoding='utf8'))
D = docx.Document(sys.argv[1] if len(sys.argv) > 1 else os.path.join(R, 'docx', 'Tabelle_Tempo_Atmosferico_v1.0.docx'))
norm = lambda s: re.sub(r'\s+', ' ', s).strip()
PAR = [norm(p.text) for p in D.paragraphs if norm(p.text)]
TAB = [[[norm(c.text) for c in r.cells] for r in t.rows] for t in D.tables]
CELLE = {c for t in TAB for r in t for c in r}
TESTI = set(PAR) | {re.sub(r'^• ', '', p) for p in PAR} | CELLE   # le voci di elenco del Word cominciano con «•»
TUTTO = '\n'.join(PAR + [c for t in TAB for r in t for c in r])
err, n = [], {'controlli': 0}
def ok(cond, msg):
    n['controlli'] += 1
    if not cond: err.append(msg)
def c_testo(t, dove):  # un testo del JSON = un paragrafo o una cella (le voci di elenco senza «• »)
    for parte in t.split('\n'): ok(parte in TESTI, f'{dove}: testo diverso dal Word: {parte[:90]}')

# --- 1. tabelle del tempo, dal Word
def intervallo(s):
    a, _, b = s.partition('-'); a = 100 if a == '00' else int(a); b = a if not b else 100 if b == '00' else int(b); return a, b
lette = {}
zona = None
corpo = D.element.body
from docx.table import Table
from docx.text.paragraph import Paragraph
for el in corpo.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p':
        t = norm(Paragraph(el, D).text); m = re.match(r'^(\d+)\. (.+)$', t)
        if m and int(m.group(1)) <= 17 and any(z['nome'] == m.group(2) for z in M['zone']): zona = m.group(1)
        if t.startswith('Capitolo 4'): zona = None
    elif tag == 'tbl' and zona:
        rows = [[norm(c.text) for c in r.cells] for r in Table(el, D).rows]
        if rows[0][0] == 'd%' and len(rows[0]) == 4:
            for col in (1, 3):
                nome = rows[0][col].split(' (')[0]
                st = next((s for s in M['stagioni'][zona] if s['nome'] == nome), None)
                ok(st is not None, f'zona {zona}: stagione «{nome}» non nel JSON')
                if st: lette[(zona, st['id'])] = [{'da': intervallo(r[col - 1])[0], 'a': intervallo(r[col - 1])[1], 'testo': r[col]} for r in rows[1:] if r[col - 1]]
righe = 0
for z, tab in M['tabelle'].items():
    for st, rows in tab.items():
        ok(lette.get((z, st)) == rows, f'tabella {z}/{st} diversa dal Word')
        k = 1
        for r in rows:
            ok(r['da'] == k and r['a'] >= r['da'], f'{z}/{st}: buco o sovrapposizione a {r}'); k = r['a'] + 1
            for v in r['testo'].split(' + '): ok(v in M['voci'] and M['voci'][v] in M['effetti'], f'{z}/{st}: «{v}» senza effetto')
            righe += 1
        ok(k == 101, f'{z}/{st}: finisce a {k - 1}, non a 00')
ok(set(lette) == {(z, s) for z, t in M['tabelle'].items() for s in t}, 'tabelle del Word e del JSON non coincidono')
so = M['sotterraneo']['righe']; k = 1
for r in so:
    ok(r['da'] == k, f'sotterraneo: buco a {r}'); k = r['a'] + 1
    ok(r['nome'] in M['effetti'], f'sotterraneo: «{r["nome"]}» senza effetto')
ok(k == 101, 'sotterraneo: non arriva a 00')
t16 = next(t for t in TAB if t[0] == ['d%', 'Condizione', 'Effetti'])
ok([[('00' if r['da'] == 100 else f"{r['da']:02d}") + ('' if r['da'] == r['a'] else '-' + ('00' if r['a'] == 100 else f"{r['a']:02d}")), r['nome'], r['testo']] for r in so] == t16[1:], 'tabella del Sotterraneo diversa dal Word')

# --- 2. effetti ed elenchi: testi uguali al Word
for nome, t in M['effetti'].items():
    if nome in {r['nome'] for r in so if r['testo'] != 'Vedi Capitolo 2.'}: ok(t in CELLE, f'condizione {nome}: testo diverso'); continue
    ok(nome in TESTI, f'effetto {nome}: titolo non nel Word'); c_testo(t, f'effetto {nome}')
for v, e in M['voci'].items(): ok(e in M['effetti'], f'voce {v} → {e} inesistente')
for z in M['zone']:
    ok(f"{z['n']}. {z['nome']}" in TESTI, f'zona {z["n"]}: nome'); c_testo(z['descrizione'], f'zona {z["n"]}')
    for p in z['particolarita']: c_testo(p, f'zona {z["n"]}, particolarità')
    for v in z['varianti']: ok(v['frase'] in z['particolarita'], f'zona {z["n"]}: variante {v["id"]} senza la sua frase')
    f = z['foraggiare']; riga = next(r for t in TAB for r in t if r[0] == f"{z['n']}. {z['nome']}")
    ok(riga == [riga[0], str(f['cd']), str(f['ogni']), f['acqua']['testo'], '☠' if f['avariato'] else ''], f'zona {z["n"]}: Foraggiare diverso dal Word {riga}')
    for nt in f['acqua']['note']: ok(nt['segno'] + ' ' + nt['testo'] in M['foraggiare']['note_frase'], f'zona {z["n"]}: nota {nt["segno"]}')
for t in M['tiro'] + [M['durate']['frase'], M['predire']['frase'], M['fiamme']['frase'], M['scala']['frase'], M['sotterraneo']['nota'],
          M['foraggiare']['frase'], M['foraggiare']['avariato_frase'], M['foraggiare']['note_frase'], M['foraggiare']['acqua_frase'], M['foraggiare']['modificatori']['nota']]:
    c_testo(t, 'frase')
for t in M['scala']['righe'] + M['acqua']['righe']: c_testo(t, 'elenco')
for k in ('sete', 'fame'): c_testo(M['fame_sete'][k]['testo'], k)
for sez in M['regole'] + [{'titolo': 'Navigazione', 'parti': M['navigazione']}]:
    ok(sez['titolo'] in TESTI, f'regole: titolo {sez["titolo"]}')
    for p in sez['parti']:
        if 'testo' in p: c_testo(p['testo'], sez['titolo'])
        if 'elenco' in p: [c_testo(x, sez['titolo']) for x in p['elenco']]
        if 'tabella' in p: ok(p['tabella'] in TAB, f'{sez["titolo"]}: tabella diversa')
for r in M['fiamme']['righe']: ok([r['fonte'], r['spegne']] in [x for t in TAB for x in t], f'fiamme: {r["fonte"]}')
for r in M['mare']: ok([r['stato'], r['testo_quando'], f"CD {r['nuotare']}"] in [x for t in TAB for x in t], f'stato del mare: {r["stato"]}')
ok(M['riassunto'] in TAB, 'Appendice B diversa')
mod = M['foraggiare']['modificatori']
ok([[r['condizione'], r['cd']] for r in mod['righe']] == next(t for t in TAB if t[0] == ['Condizione', 'CD'])[1:], 'modificatori di Foraggiare diversi')

# --- 3. regole del motore e loro frasi
dur = M['durate']['frase']
for v in M['durate']['lunghe']:
    nome = v.replace('Tempesta di ', '') if v.startswith('Tempesta di ') and v != 'Tempesta di vento' else v
    ok(nome in dur, f'durata 2d4: «{v}» non nella frase')
ok('2d4 ore' in dur and 'La Nebbia si dirada 1d4 ore dopo l\'alba' in dur and 'prima riga della tabella' in dur, 'durate: frase cambiata')
ok("vale dalle 10 alle 18, nelle altre ore vale il Clima caldo" in M['effetti']['Arsura / Gelo'], 'Arsura: ore cambiate')
ok(M['zone'][9]['particolarita'][0].startswith(M['notti_fredde']['frase']) and 'Le notti sono fredde come nel Deserto sabbioso.' in M['zone'][10]['particolarita'], 'notti fredde')
ok('Sopravvivenza CD 15, altre 24 ore ogni 5 punti oltre la CD' in M['predire']['frase'], 'Predire il tempo cambiato')
fuoco = next(x for x in M['regole'][2]['parti'][0]['elenco'] if x.startswith('Fuoco.'))
for z, nome in M['senza_legna'].items(): ok(nome in fuoco, f'zone senza legna: {nome}')
for v in mod['piu2']: ok(v in M['voci'], f'Foraggiare +2: «{v}» non è una voce')
for v in mod['piu5']: ok(v in M['voci'], f'Foraggiare +5: «{v}» non è una voce')
ok('Inverno (zone a quattro stagioni), Autunno-inverno (Foreste boreali e Tundra), Stagione secca (Savana e steppa)' in mod['righe'][0]['condizione'], 'stagione povera cambiata')
a = '\n'.join(M['acqua']['righe'])
ok('Ne servono 2 con Arsura o Afa, 3 con Canicola.' in a and M['acqua']['moltiplicatori'] == {'Arsura': 2, 'Afa': 2, 'Canicola': 3}, 'moltiplicatori dell\'acqua')
ok('con Pioggia battente, Acquazzone, Monsone o Temporale chi ha recipienti raccoglie la razione del giorno' in a, 'pioggia')
ok('10 minuti al fuoco per razione' in a and '12 litri, cioè 3 razioni' in a and 'Razione d\'acqua: 4 litri al giorno' in a, 'bollire, tavoletta, litri')
ok('quarto giorno' in M['fame_sete']['sete']['testo'] and 'settimo giorno' in M['fame_sete']['fame']['testo'] and 'Costituzione CD 10' in M['fame_sete']['fame']['testo'], 'Fame e sete: giorni o CD cambiati')
ok('Tiro Salvezza su Tempra CD 12' in M['foraggiare']['avariato_frase'] and 'di 5 o più' in M['foraggiare']['avariato_frase'] and 'nauseato per 1d4 ore' in M['foraggiare']['avariato_frase'], 'cibo avariato cambiato')
ok('2 ore fuori dalla marcia' in M['foraggiare']['frase'] and 'Professione (pescatore)' in M['foraggiare']['frase'], 'Foraggiare: frase cambiata')
for z in M['zone']:
    st = [s['id'] for s in M['stagioni'][z['id']]]
    ok(z['id'] == '16' and not st or len(st) in (2, 4), f'zona {z["id"]}: stagioni {st}')
for vecchia, nuova in M['migrazione'].items(): ok(nuova in M['tabelle'], f'migrazione {vecchia} → {nuova}')

for e in err: print('ERRORE', e)
print(f"zone {len(M['zone'])}, tabelle {len(lette)}, righe {righe}, effetti {len(M['effetti'])}, sotterraneo {len(so)}, controlli {n['controlli']}, errori {len(err)}")
sys.exit(1 if err else 0)
